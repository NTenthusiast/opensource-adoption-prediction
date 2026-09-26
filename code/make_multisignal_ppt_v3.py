"""根据近期扩样回测的 CSV 生成可编辑组会 PPT。"""
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

from make_multisignal_ppt import add_text, slide, table, NAVY, TEAL, ORANGE


BASE = Path(__file__).resolve().parents[1]
TABLES = BASE / "code" / "results" / "tables"
OUT = BASE / "组会汇报_2026-09-26" / "组会汇报_三维动态预警升级版.pptx"
RULES = [
    ("原问卷规则", "原规则"),
    ("只看评论", "评论规则"),
    ("双源同时收缩", "双源同时收缩"),
    ("任一信号收缩", "任一信号收缩"),
]


def result_row(frame, outcome, prefix):
    return frame[(frame["因变量"] == outcome) & frame["规则"].str.startswith(prefix)].iloc[0]


def result_table(s, comparison, outcome):
    body = [["规则", "警报", "命中", "误报", "漏报", "正确未报", "命中率", "找回率"]]
    for label, prefix in RULES:
        r = result_row(comparison, outcome, prefix)
        tn = int(r["样本"] - r["命中"] - r["误报"] - r["漏报"])
        body.append([label, int(r["警报数"]), int(r["命中"]), int(r["误报"]),
                     int(r["漏报"]), tn, f"{r['精确率']:.1%}", f"{r['召回率']:.1%}"])
    table(s, body, 0.62, 1.55, [2.35, 1.0, 1.0, 1.0, 1.0, 1.6, 1.85, 1.85], 0.71, 15)


def number_or_dash(value, alarms):
    return f"{value:.1%}" if alarms else "—"


def main():
    comparison = pd.read_csv(TABLES / "multisignal_recent_comparison.csv", encoding="utf-8-sig")
    by_year = pd.read_csv(TABLES / "multisignal_recent_by_year.csv", encoding="utf-8-sig")
    cases = pd.read_csv(TABLES / "multisignal_recent_cases.csv", encoding="utf-8-sig")
    coverage = pd.read_csv(TABLES / "multisignal_recent_coverage.csv", encoding="utf-8-sig")
    latest = pd.read_csv(TABLES / "multisignal_recent_latest.csv", encoding="utf-8-sig")
    stable = pd.read_csv(TABLES / "multisignal_recent_pre2025_comparison.csv", encoding="utf-8-sig")
    shift = pd.read_csv(TABLES / "multisignal_recent_survey_shift.csv", encoding="utf-8-sig")
    survey_coverage = pd.read_csv(BASE / "code" / "data" / "processed" / "survey_coverage_recent.csv",
                                  encoding="utf-8-sig")
    so_name = "下一年问卷使用率下降"
    gh_name = "下一年主仓评论量下降"
    so = cases[cases.future_so_ok]
    gh = cases[cases.future_gh_ok]
    so_old = result_row(comparison, so_name, "原规则")
    so_dual = result_row(comparison, so_name, "双源同时收缩")
    gh_old = result_row(comparison, gh_name, "原规则")
    gh_dual = result_row(comparison, gh_name, "双源同时收缩")
    stable_old = result_row(stable, so_name, "原规则")
    stable_dual = result_row(stable, so_name, "双源同时收缩")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY
    add_text(s, "开源技术三维动态预警", 0.85, 1.35, 11.7, 1.0, 43,
             RGBColor(255, 255, 255), True)
    add_text(s, "扩大技术样本后的近期年度回测", 0.88, 2.58, 11.4, 0.68,
             25, RGBColor(214, 238, 237))
    add_text(s, "组会汇报", 0.89, 6.18, 9.0, 0.4, 17, RGBColor(214, 238, 237))

    s = slide(prs, "本轮只改样本与年份", "对照：前版结果 CSV；本轮结果 CSV 与官方调查原始数据")
    table(s, [
        ["", "前版", "本轮"],
        ["可评估技术", "5 项", f"{len(coverage)} 项"],
        ["使用率结果记录", "12 条", f"{len(so)} 条"],
        ["评论量结果记录", "17 条", f"{len(gh)} 条"],
        ["预警年", "较早的连续年份", "2022—2024 年"],
        ["最终结果年", "截至 2023 年", "截至 2025 年"],
    ], 0.71, 1.43, [3.2, 4.35, 4.75], 0.69, 18)
    add_text(s, "两版的规则和命中定义相同；不同样本的百分比不可当作模型提升直接相减。",
             0.81, 6.24, 11.8, 0.57, 18, ORANGE)

    s = slide(prs, "规则保持不变", "问卷：Stack Overflow；评论：GitHub 主仓 Issue/PR 对话评论")
    table(s, [
        ["维度", "当年可用数据", "作用"],
        ["使用规模", "问卷使用率", "同级警报中安排核查顺序"],
        ["使用意愿", "想用率－使用率连续两年为负", "问卷警报"],
        ["社区活动", "当年评论量比上年少", "评论警报"],
    ], 0.72, 1.51, [2.35, 5.0, 4.95], 0.88, 17)
    add_text(s, "双源＝两种警报同时出现；任一信号＝出现其中一种就警报。",
             0.83, 5.56, 11.6, 0.72, 20, TEAL)
    add_text(s, "评论量是讨论数量，不代表满意度或技术整体用户规模。",
             0.83, 6.17, 11.6, 0.62, 18, ORANGE)

    s = slide(prs, "怎么检验：先警报，再看下一年", "源数据：panel_extended.csv；github_comment_yearly.csv")
    add_text(s, "① 用 t－1 年和 t 年的问卷、评论量，决定 t 年是否报警。",
             0.83, 1.48, 11.7, 0.8, 23)
    add_text(s, "② 用 t＋1 年问卷使用率或评论量，核对当年的警报。",
             0.83, 2.74, 11.7, 0.8, 23)
    add_text(s, "③ 逐年记录命中、误报、漏报和正确未报；四种规则用同一批有效记录。",
             0.83, 4.02, 11.7, 1.1, 22, TEAL)
    add_text(s, "2025 年只作已完成的结果年；不把尚未结束的 2026 年当作回测结果。",
             0.83, 5.83, 11.7, 0.86, 18, ORANGE)

    s = slide(prs, "近期样本到底有多少", "完整技术与仓库清单：multisignal_recent_coverage.csv")
    rows = [["预警年", "使用率可评估", "评论量可评估"]]
    for year in (2022, 2023, 2024):
        rows.append([year, int((so.year == year).sum()), int((gh.year == year).sum())])
    rows.append(["合计", len(so), len(gh)])
    table(s, rows, 0.84, 1.52, [3.3, 4.3, 4.3], 0.76, 19)
    add_text(s, f"纳入 {len(coverage)} 项技术；一项技术在不同预警年可贡献多条记录。",
             0.87, 5.84, 11.6, 0.77, 21, TEAL)
    add_text(s, "只有前一年、当年和下一年所需数据都齐全时，才纳入对应终点。",
             0.87, 6.40, 11.6, 0.48, 16, ORANGE)

    s = slide(prs, "测试一：下一年使用率是否下降", f"{len(so)} 条技术年度记录；其中 {int(so.usage_down.sum())} 条实际下降")
    result_table(s, comparison, so_name)
    add_text(s, f"双源：{int(so_dual['警报数'])} 次警报命中 {int(so_dual['命中'])} 次；旧规则：{int(so_old['警报数'])} 次中命中 {int(so_old['命中'])} 次。",
             0.74, 5.53, 11.75, 0.62, 20, TEAL)
    add_text(s, f"双源漏报 {int(so_dual['漏报'])} 次，旧规则漏报 {int(so_old['漏报'])} 次；同时看命中率和找回率。",
             0.74, 6.20, 11.75, 0.58, 19, ORANGE)

    s = slide(prs, "测试二：下一年评论量是否下降", f"{len(gh)} 条技术年度记录；其中 {int(gh.comments_down.sum())} 条实际下降")
    result_table(s, comparison, gh_name)
    add_text(s, f"双源误报 {int(gh_dual['误报'])} 次，旧规则误报 {int(gh_old['误报'])} 次。",
             0.74, 5.53, 11.75, 0.62, 20, TEAL)
    add_text(s, f"双源漏报 {int(gh_dual['漏报'])} 次，旧规则漏报 {int(gh_old['漏报'])} 次。",
             0.74, 6.20, 11.75, 0.58, 19, ORANGE)

    s = slide(prs, "分开按年看，避免合并结果掩盖波动", "下一年使用率结果；逐年明细：multisignal_recent_by_year.csv")
    rows = [["预警年", "样本", "实际下降", "旧规则命中率", "旧规则找回率", "双源命中率", "双源找回率"]]
    for year in (2022, 2023, 2024):
        a = by_year[(by_year.因变量 == so_name) & by_year.规则.str.startswith("原规则") & (by_year.预警年 == year)].iloc[0]
        b = by_year[(by_year.因变量 == so_name) & by_year.规则.str.startswith("双源同时收缩") & (by_year.预警年 == year)].iloc[0]
        actual = int(a["命中"] + a["漏报"])
        rows.append([year, int(a["样本"]), actual,
                     number_or_dash(a["精确率"], a["警报数"]),
                     number_or_dash(a["召回率"], actual),
                     number_or_dash(b["精确率"], b["警报数"]),
                     number_or_dash(b["召回率"], actual)])
    table(s, rows, 0.71, 1.55, [1.45, 1.0, 1.45, 2.05, 2.05, 2.05, 2.05], 0.86, 15)
    add_text(s, "2024 年这组技术没有下一年下降；找回率的分母为零，故显示“—”。",
             0.82, 5.53, 11.6, 0.92, 19, ORANGE)

    s = slide(prs, "看几条真实记录如何判定", "示例取自 2024 年预警；2025 年问卷结果只用于事后核验")
    examples = so[so.year == 2024].sort_values(["alert_score", "usage"], ascending=False).head(6)
    rows = [["技术", "当年问卷", "当年评论", "当年输出", "后续使用率"]]
    for _, item in examples.iterrows():
        rows.append([item.tech, "警报" if item.so_alert else "无",
                     "警报" if item.gh_alert else "无", item["状态"],
                     "下降" if item.usage_down else "未下降"])
    table(s, rows, 0.73, 1.5, [2.4, 1.7, 1.7, 3.8, 2.7], 0.67, 16)
    add_text(s, "本页仅示例决策过程；总体结论以全部记录的回测表为准。",
             0.83, 6.41, 11.5, 0.46, 17, TEAL)

    s = slide(prs, "当前可用的核查清单", "依据 2025 年完整问卷与主仓评论；2026 年结果尚未用于回测")
    rows = [["技术", "问卷", "评论", "输出", "使用率"]]
    for _, item in latest.head(7).iterrows():
        rows.append([item.tech, "警报" if item.so_alert else "无",
                     "警报" if item.gh_alert else "无", item["状态"],
                     f"{item.usage:.1%}"])
    table(s, rows, 0.73, 1.37, [2.25, 1.6, 1.6, 4.55, 2.3], 0.62, 16)
    add_text(s, f"共 {len(latest)} 项有两路历史；此处展示优先核查的前 7 项。",
             0.82, 6.54, 11.6, 0.35, 16, TEAL)

    s = slide(prs, "必须单看 2025 年问卷的整体上移", "诊断表：multisignal_recent_survey_shift.csv；survey_coverage_recent.csv")
    all_up = int(shift["上升技术数"].sum())
    all_common = int(shift["共同技术数"].sum())
    web = shift[shift.category == "Webframework"].iloc[0]
    lang = survey_coverage[survey_coverage.category == "Language"].set_index("year")
    add_text(s, f"在两年都出现的 {all_common} 项技术中，{all_up} 项的问卷使用率上升；",
             0.81, 1.43, 11.6, 0.8, 23, TEAL, True)
    add_text(s, f"Web 框架为 {int(web['上升技术数'])}/{int(web['共同技术数'])}；本轮 2024 年预警对应的 {len(so[so.year == 2024])} 项也全部上升。",
             0.81, 2.44, 11.6, 0.85, 22)
    add_text(s, f"语言使用题有效回答：2024 年 {int(lang.loc[2024, '使用题有效回答数']):,} 人，"
                f"2025 年 {int(lang.loc[2025, '使用题有效回答数']):,} 人。",
             0.81, 3.72, 11.6, 0.83, 21)
    add_text(s, f"只看预警年 2022—2023：{int(stable_old['样本'])} 条记录；"
                f"双源命中率 {stable_dual['精确率']:.1%}，旧规则 {stable_old['精确率']:.1%}。",
             0.81, 4.95, 11.6, 1.0, 20, TEAL)
    add_text(s, "这提示调查受访者或选项口径可能变化；不能把整体上移直接解释为技术都增长。",
             0.81, 6.18, 11.6, 0.65, 18, ORANGE)

    s = slide(prs, "怎样理解这次扩样", "方法与筛选限制见研究记录_多源因变量与动态预警.md")
    add_text(s, "样本增加后，先看命中和漏报是否仍能同时接受。",
             0.83, 1.53, 11.55, 0.77, 24, TEAL, True)
    add_text(s, "不同年份的问卷受访者不是同一批人；技术选项与有效回答数也会变化。",
             0.83, 2.79, 11.55, 0.9, 21)
    add_text(s, "一个 GitHub 主仓只代表部分社区活动；评论减少不一定意味着技术衰退。",
             0.83, 4.13, 11.55, 1.04, 21)
    add_text(s, "仓库按可核验的完整年度计数纳入，不能代表所有开源技术。",
             0.83, 5.67, 11.55, 0.9, 20, ORANGE)

    s = slide(prs, "结论：扩大样本后保留有条件的预警用途", "逐项结果、仓库和年份均可在本项目 CSV 中复核")
    add_text(s, f"使用率下降检验：双源命中率 {so_dual['精确率']:.1%}，找回率 {so_dual['召回率']:.1%}。",
             0.83, 1.51, 11.7, 0.83, 24, TEAL, True)
    add_text(s, f"原问卷规则：命中率 {so_old['精确率']:.1%}，找回率 {so_old['召回率']:.1%}。",
             0.83, 2.78, 11.7, 0.8, 23)
    add_text(s, "双源可用于优先核查的候选；2025 年问卷整体上移使未来稳定性仍待验证。",
             0.83, 4.25, 11.55, 1.02, 22)
    add_text(s, "扩大样本提升了观察范围，不能自动消除样本选择和跨平台口径问题。",
             0.83, 5.77, 11.55, 0.82, 19, ORANGE)

    OUT.parent.mkdir(exist_ok=True)
    prs.save(OUT)
    print(f"已生成 {OUT}（{len(prs.slides)} 页）")


if __name__ == "__main__":
    main()
