"""根据可复现结果生成本次组会汇报（图表保持 PowerPoint 可编辑）。"""
from pathlib import Path
import shutil

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parents[1]
SRC = Path(r"C:\Users\83786\Downloads\8.26.pptx")
DEST = BASE / "组会汇报_2026-09-26"
OUT = DEST / "组会汇报_多源因变量与动态预警_无汇报日期.pptx"
TBL = BASE / "code" / "results" / "tables"
DATA = BASE / "code" / "data" / "processed"

NAVY = RGBColor(22, 42, 65)
TEAL = RGBColor(0, 126, 128)
ORANGE = RGBColor(195, 102, 35)
GREY = RGBColor(93, 104, 114)
LIGHT = RGBColor(248, 250, 251)
FONT = "Microsoft YaHei"


def add_text(slide, text, x, y, w, h, size=20, color=NAVY, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.space_after = Pt(10)
        if align is not None:
            p.alignment = align
    return box


def slide(prs, title, source=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = LIGHT
    add_text(s, title, 0.65, 0.35, 12, 0.75, 31, NAVY, True)
    if source:
        add_text(s, source, 0.67, 7.08, 11.9, 0.22, 10, GREY)
    add_text(s, str(len(prs.slides)).zfill(2), 12.35, 7.08, 0.45, 0.22, 10, GREY, align=PP_ALIGN.RIGHT)
    return s


def table(slide, rows, x, y, widths, row_h=0.52, font_size=16):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y),
                                   Inches(sum(widths)), Inches(row_h * len(rows)))
    t = shape.table
    for j, width in enumerate(widths):
        t.columns[j].width = Inches(width)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, value in enumerate(row):
            c = t.cell(i, j)
            c.text = str(value)
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if i == 0 else (RGBColor(234, 241, 243) if i % 2 else LIGHT)
            for p in c.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(font_size)
                p.font.bold = i == 0
                p.font.color.rgb = RGBColor(255, 255, 255) if i == 0 else NAVY
    return shape


def main():
    so = pd.read_csv(TBL / "warning_validation_so.csv", encoding="utf-8-sig")
    by_year = pd.read_csv(TBL / "warning_validation_by_year.csv", encoding="utf-8-sig")
    gh = pd.read_csv(TBL / "warning_validation_github.csv", encoding="utf-8-sig").iloc[0]
    cases = pd.read_csv(TBL / "warning_validation_github_cases.csv", encoding="utf-8-sig")
    yearly = pd.read_csv(DATA / "github_comment_yearly.csv", encoding="utf-8-sig")
    latest = pd.read_csv(TBL / "multisignal_warning_latest.csv", encoding="utf-8-sig")
    DEST.mkdir(exist_ok=True)
    shutil.copy2(SRC, DEST / "8.26_上次组会汇报.pptx")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY
    add_text(s, "多源因变量与动态预警", 0.85, 1.45, 11.6, 1.0, 43, RGBColor(255,255,255), True)
    add_text(s, "开源技术采纳研究的前瞻检验与 GitHub 评论对照", 0.88, 2.55, 11.4, 0.65,
             24, RGBColor(214, 238, 237))
    add_text(s, "组会汇报", 0.89, 6.18, 9.0, 0.4,
             17, RGBColor(214, 238, 237))

    s = slide(prs, "这次研究要回答的问题", "资料：前期汇报与跨平台、评论状态研究总结")
    add_text(s, "研究问题：把“用户评论”等问卷外结果变量纳入实验。", 0.78, 1.45, 11.8, 0.65, 26, TEAL, True)
    add_text(s, "本次新增两个检验：\n① 旧预警能否预测下一年使用率、想用率、欣赏率下降？\n② 同一信号能否对应下一年 GitHub 主仓评论量下降？",
             0.85, 2.38, 11.6, 2.4, 22)
    add_text(s, "已有研究提示：跨平台“水平对应”不等于时间领先；评论状态在小样本中未稳定改善预测。",
             0.85, 5.44, 11.3, 0.9, 19, ORANGE)

    s = slide(prs, "先修正上次汇报的解释边界", "依据：本项目数据构造代码；Stack Overflow 2022 调查方法说明")
    table(s, [
        ["上次用语", "本次准确口径"],
        ["β₁≈0.80 是用户留存率", "跨年技术比例的回归系数；无法跟踪同一用户"],
        ["β₂≈0.24 是意愿转化率", "想用比例的预测关联；不是个人转化概率"],
        ["负渴望缺口是净流出", "同年两道题的比例差；需用未来结果验证"],
    ], 0.75, 1.62, [4.0, 8.0], 0.88, 17)
    add_text(s, "因此，以下所有结果都称“预警关联”，不称因果效应。", 0.84, 5.75, 11.7, 0.65, 22, TEAL, True)

    s = slide(prs, "数据与时间顺序", "来源：项目 panel.csv；GitHub REST API /repos/{owner}/{repo}/issues/comments")
    add_text(s, "Stack Overflow：2018—2022 年技术面板。2019—2021 年发出警报，检查下一年问卷结果。",
             0.81, 1.55, 11.7, 1.1, 22)
    add_text(s, "GitHub：5 个公开主仓，2018—2023 年逐年统计 Issue/PR 对话评论。用 t 年警报检查 t+1 年评论变化。",
             0.81, 3.01, 11.7, 1.2, 22)
    add_text(s, "评论指标是“参与数量”，不代表情绪、满意度或全部技术生态；API 完整分页后计数，不保存正文。",
             0.81, 5.13, 11.5, 1.0, 20, ORANGE)

    s = slide(prs, "问卷内部：警报与下一年使用率下降有关", "来源：code/results/tables/warning_validation_so.csv；按技术聚类重采样 2,000 次")
    u = so.iloc[0]
    data = CategoryChartData()
    data.categories = ["发出警报", "未发警报"]
    data.add_series("下一年使用率下降比例", [100*u["精确率"], 100*u["无警报下降率"]])
    chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.98), Inches(1.58),
                               Inches(6.55), Inches(4.7), data).chart
    chart.has_legend = False
    chart.value_axis.maximum_scale = 100
    chart.value_axis.minimum_scale = 0
    chart.plots[0].has_data_labels = True
    chart.plots[0].data_labels.number_format = '0.0"%"'
    chart.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    add_text(s, f"{int(u['样本'])} 个有效样本\n{int(u['警报数'])} 次警报，{int(u['警报后下降'])} 次随后下降",
             8.1, 1.83, 4.5, 1.6, 21, NAVY, True)
    add_text(s, f"风险差 +{100*u['风险差']:.1f} 个百分点\n聚类区间 [{100*u['区间下限']:.1f}, {100*u['区间上限']:.1f}]",
             8.1, 4.04, 4.5, 1.4, 20, TEAL)

    s = slide(prs, "换一个因变量，预警表现随之变化", "来源：code/results/tables/warning_validation_so.csv")
    metric_names = ["使用率", "想用率", "欣赏率"]
    rows = [["下一年下降", "警报组", "无警报组", "差值（百分点）"]]
    for name, (_, r) in zip(metric_names, so.iterrows()):
        rows.append([name, f"{100*r['精确率']:.1f}%", f"{100*r['无警报下降率']:.1f}%",
                     f"{100*r['风险差']:+.1f}"])
    table(s, rows, 1.1, 1.62, [3.2, 2.6, 2.8, 3.2], 0.83, 18)
    add_text(s, "预警关联集中在“下一年使用率”。欣赏率方向相反，不能把三项因变量合称为同一类衰退。",
             1.12, 5.43, 11.0, 1.0, 21, ORANGE)

    s = slide(prs, "时间稳定性：2019 年警报没有领先优势", "来源：code/results/tables/warning_validation_by_year.csv")
    rows = [["预警年份", "样本数", "警报后使用率下降", "无警报使用率下降"]]
    for _, r in by_year.iterrows():
        rows.append([int(r["year"]), int(r["样本"]),
                     f"{100*r['警报后usage下降率']:.1f}%",
                     f"{100*r['无警报usage下降率']:.1f}%"])
    table(s, rows, 0.95, 1.62, [2.4, 2.0, 3.6, 3.8], 0.83, 17)
    add_text(s, "合并样本的优势主要来自 2020、2021 年；不能宣称已得到稳定的跨年预警规则。",
             1.0, 5.46, 11.3, 0.95, 21, ORANGE)

    s = slide(prs, "GitHub 评论：真实的问卷外结果变量", "来源：GitHub REST API；code/data/processed/github_comment_yearly.csv")
    years = sorted(yearly.year.unique())
    data = CategoryChartData(); data.categories = [str(y) for y in years]
    for tech, g in yearly.groupby("tech"):
        data.add_series(tech, [int(g.loc[g.year == y, "issue_pr_comments"].iloc[0]) for y in years])
    chart = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.72), Inches(1.55),
                               Inches(8.65), Inches(4.95), data).chart
    chart.has_legend = True; chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    add_text(s, "按创建时间统计。\n只含主仓 Issue/PR 对话评论。\n绝对规模不宜跨技术直接比较。",
             9.65, 1.91, 2.83, 3.3, 17)

    s = slide(prs, "跨平台：评论量检验仍是小样本探索", "来源：code/results/tables/warning_validation_github.csv 及逐项明细")
    add_text(s, f"{int(gh['样本'])} 个可比“技术×年份”观测", 0.9, 1.6, 11.5, 0.6, 27, TEAL, True)
    add_text(s, f"{int(gh['警报数'])} 次问卷警报，其中 {int(gh['警报后下降'])} 次下一年主仓评论量下降。",
             0.9, 2.63, 11.5, 0.7, 22)
    add_text(s, f"警报组下降率 {100*gh['精确率']:.1f}%；无警报组 {100*gh['无警报下降率']:.1f}%。",
             0.9, 3.71, 11.5, 0.7, 22)
    add_text(s, "结论仅适用于这批仓库。项目治理、版本发布、机器人及仓库迁移均可能改变评论量。",
             0.9, 5.28, 11.5, 1.0, 20, ORANGE)

    s = slide(prs, "双源信号用于核查优先级", "来源：code/results/tables/multisignal_warning_latest.csv；观察截点 2022 年")
    rows = [["技术", "问卷警报", "评论动量(log)", "核查标签"]]
    for _, r in latest.iterrows():
        rows.append([r["tech"], "不可判定" if not r["so_history_ok"] else ("是" if r["warning"] else "否"),
                     f"{r['comment_momentum']:+.2f}", r["信号组合"]])
    table(s, rows, 0.84, 1.49, [2.3, 2.2, 2.2, 5.2], 0.7, 16)
    add_text(s, "同向收缩：优先人工核查。信号冲突：检查仓库代表性与问卷口径。标签不是风险概率。",
             0.87, 6.06, 11.6, 0.62, 18, TEAL)

    s = slide(prs, "研究结论与下一轮实验", "依据：本次回测及跨平台、评论状态研究")
    add_text(s, "可支持：旧渴望缺口规则与下一年问卷使用率下降有关。", 0.81, 1.5, 11.9, 0.75, 23, TEAL, True)
    add_text(s, "尚不支持：把该规则直接解释为用户流失、情绪下降，或跨平台通用预警。",
             0.81, 2.66, 11.9, 0.9, 22, ORANGE)
    add_text(s, "下一轮优先做：扩大独立仓库样本；区分人工/机器人评论及 Issue/PR；统一问卷技术别名；用未见过的年份检验。",
             0.81, 4.38, 11.8, 1.6, 21)

    prs.save(OUT)
    print(f"已生成：{OUT}（{len(prs.slides)} 页）")


if __name__ == "__main__":
    main()
