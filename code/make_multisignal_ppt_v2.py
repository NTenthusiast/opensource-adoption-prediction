"""生成三维动态预警升级版组会 PPT，所有结果表格均可编辑。"""
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.util import Inches

from make_multisignal_ppt import add_text, slide, table, NAVY, TEAL, ORANGE

BASE = Path(__file__).resolve().parents[1]
TABLES = BASE / "code" / "results" / "tables"
OUT = BASE / "组会汇报_2026-09-26" / "组会汇报_三维动态预警升级版.pptx"


def selected(rows, outcome, rule):
    return rows[(rows["因变量"] == outcome) & rows["规则"].str.startswith(rule)].iloc[0]


def rule_table(s, comparison, outcome):
    labels = [
        ("原规则", "原规则"),
        ("只看评论", "评论规则"),
        ("双源同时收缩", "双源同时收缩"),
        ("任一信号收缩", "任一信号收缩"),
    ]
    body = [["规则", "警报", "命中", "误报", "漏报", "正确未报", "命中率", "找回率"]]
    for label, prefix in labels:
        r = selected(comparison, outcome, prefix)
        correct_no_alert = int(r["样本"] - r["命中"] - r["误报"] - r["漏报"])
        body.append([label, int(r["警报数"]), int(r["命中"]), int(r["误报"]),
                     int(r["漏报"]), correct_no_alert,
                     f"{100*r['精确率']:.1f}%", f"{100*r['召回率']:.1f}%"])
    table(s, body, 0.62, 1.54, [2.35, 1.0, 1.0, 1.0, 1.0, 1.6, 1.85, 1.85], 0.71, 15)


def main():
    comparison = pd.read_csv(TABLES / "multisignal_backtest_comparison.csv", encoding="utf-8-sig")
    latest = pd.read_csv(TABLES / "multisignal_warning_v2_latest.csv", encoding="utf-8-sig")
    cases = pd.read_csv(TABLES / "multisignal_backtest_cases.csv", encoding="utf-8-sig")
    so = pd.read_csv(TABLES / "warning_validation_so.csv", encoding="utf-8-sig")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY
    add_text(s, "开源技术三维动态预警", 0.85, 1.35, 11.7, 1.0, 43,
             RGBColor(255, 255, 255), True)
    add_text(s, "加入 GitHub 评论活动后的规则与回测", 0.88, 2.58, 11.4, 0.68,
             25, RGBColor(214, 238, 237))
    add_text(s, "组会汇报", 0.89, 6.18, 9.0, 0.4, 17, RGBColor(214, 238, 237))

    s = slide(prs, "为什么改预警系统", "依据：项目原始预警代码及本次双源回测")
    add_text(s, "原系统的使用率、想用率、欣赏率都来自同一份调查。", 0.82, 1.48, 11.6, 0.9,
             25, NAVY, True)
    add_text(s, "新增独立的 GitHub 主仓评论量，观察社区讨论是否也在收缩。",
             0.82, 2.88, 11.5, 0.9, 24, TEAL)
    add_text(s, "目标：把预警变成可操作的核查顺序，并检验它是否比旧规则更准。",
             0.82, 4.57, 11.6, 1.0, 23)

    s = slide(prs, "新版三维分别看什么", "问卷指标来自 panel.csv；评论数来自 GitHub 公开 Issue/PR 对话评论")
    table(s, [
        ["维度", "具体指标", "在预警中的用途"],
        ["使用规模", "当年问卷使用率", "同级警报中优先核查影响面较大的技术"],
        ["使用意愿", "想用率－使用率连续两年为负", "发出问卷警报"],
        ["社区活动", "主仓评论数较上一年下降", "发出评论警报"],
    ], 0.71, 1.56, [2.4, 4.65, 5.25], 0.91, 17)
    add_text(s, "评论量衡量讨论数量，不等于满意度；使用率用于排序，不冒充衰退概率。",
             0.82, 5.68, 11.6, 0.7, 20, ORANGE)

    s = slide(prs, "双源预警怎样发出", "阈值在回测前固定；只使用当年及以前的信息")
    table(s, [
        ["问卷警报", "评论警报", "输出", "核查动作"],
        ["有", "有", "双源收缩", "优先人工核查"],
        ["有", "无", "问卷单源警报", "检查问卷口径与仓库代表性"],
        ["无", "有", "评论单源警报", "检查项目维护及讨论变化"],
        ["无", "无", "暂无收缩信号", "继续观察"],
    ], 0.65, 1.43, [2.1, 2.1, 2.65, 5.45], 0.78, 16)
    add_text(s, "历史年份不连续时输出“历史不足”，不把缺数据算成安全。",
             0.77, 6.11, 11.8, 0.53, 18, TEAL)

    s = slide(prs, "数据与检验范围", "来源：Stack Overflow 技术面板；GitHub 公开评论 API；逐项明细见结果 CSV")
    add_text(s, "7 个公开主仓的年度评论数", 0.83, 1.53, 11.5, 0.6, 27, TEAL, True)
    add_text(s, "其中 5 项技术有连续问卷历史，可比较旧规则和新规则。",
             0.83, 2.49, 11.5, 0.8, 22)
    add_text(s, "使用率检验：12 条＝Kotlin、Ruby 各 3 条；Express、Flask、jQuery 各 2 条。\n评论量检验：17 条＝Kotlin、Ruby 各 4 条；Express、Flask、jQuery 各 3 条。",
             0.83, 3.68, 11.65, 1.25, 18)
    add_text(s, "评论量多 5 条：这 5 项技术的 2022 年记录可核对 2023 年评论，缺少对应问卷。",
             0.83, 5.11, 11.65, 0.65, 17, TEAL)
    add_text(s, "一项技术每个可检验年份算一条；样本少，百分比只作探索性比较。",
             0.83, 6.02, 11.5, 0.55, 18, ORANGE)

    s = slide(prs, "旧规则有一定信号，但并不稳定", "来源：code/results/tables/warning_validation_so.csv 及 warning_validation_by_year.csv")
    base = so.iloc[0]
    table(s, [
        ["", "下一年使用率下降", "没有下降", "下降比例"],
        ["发出问卷警报", 46, 23, "66.7%"],
        ["没有问卷警报", 33, 42, "44.0%"],
    ], 1.0, 1.66, [3.1, 3.35, 2.8, 2.4], 0.87, 18)
    add_text(s, f"共 {int(base['样本'])} 条记录。警报组的下降比例高出 22.7 个百分点。",
             1.05, 4.65, 11.2, 0.75, 21, TEAL)
    add_text(s, "按年份分开看，结果有反转；旧规则还不能作为稳定的自动判定。",
             1.05, 5.69, 11.3, 0.73, 19, ORANGE)

    s = slide(prs, "测试一：能否筛出下一年使用率下降", "五项技术、12 条可比记录；每行满足：命中＋误报＋漏报＋正确未报＝12")
    rule_table(s, comparison, "下一年问卷使用率下降")
    add_text(s, "双源同时收缩：3 次警报都命中；旧规则 5 次中命中 4 次。",
             0.74, 5.49, 11.75, 0.58, 20, TEAL)
    add_text(s, "代价：双源规则漏掉 4 次下降，旧规则漏掉 3 次。",
             0.74, 6.15, 11.7, 0.53, 19, ORANGE)

    s = slide(prs, "命中更准，也可能漏掉更多", "来源：multisignal_backtest_comparison.csv；该图只比较相同的 12 条记录")
    d = CategoryChartData()
    d.categories = ["旧问卷规则", "双源同时收缩"]
    d.add_series("命中率", [80.0, 100.0])
    d.add_series("找回率", [57.1, 42.9])
    chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.85), Inches(1.53),
                               Inches(7.3), Inches(4.76), d).chart
    chart.has_legend = False
    chart.value_axis.minimum_scale = 0; chart.value_axis.maximum_scale = 100
    chart.plots[0].has_data_labels = True
    chart.plots[0].data_labels.number_format = '0.0"%"'
    chart.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    add_text(s, "蓝色：命中率，警报有多准。\n红色：找回率，下降检出多少。",
             8.38, 1.81, 4.1, 2.0, 18)
    add_text(s, "双源规则适合排人工核查的优先级，尚不能替代广覆盖监测。",
             8.38, 4.63, 4.15, 1.45, 18, ORANGE)

    s = slide(prs, "测试二：能否筛出下一年评论量下降", "五项技术、17 条可比记录；每行满足：命中＋误报＋漏报＋正确未报＝17")
    rule_table(s, comparison, "下一年主仓评论量下降")
    add_text(s, "双源规则误报更少（1 次对 3 次），但漏报更多（6 次对 4 次）。",
             0.74, 5.49, 11.75, 0.58, 20, TEAL)
    add_text(s, "只要任一信号收缩就警报，可找回 8/10 次下降，但会发出 4 次误报。",
             0.74, 6.15, 11.7, 0.53, 19, ORANGE)

    s = slide(prs, "具体技术说明了信号冲突的价值", "表中后续评论数只用于事后检验，没有进入当年的预警")
    chosen = cases[(cases["year"] == 2022) & cases["tech"].isin(["Express", "Flask", "Ruby", "Kotlin", "jQuery"])].copy()
    ordered = ["Express", "Flask", "Ruby", "Kotlin", "jQuery"]
    rows = [["技术", "当年信号", "当年评论", "后续评论", "实际变化"]]
    for tech in ordered:
        r = chosen[chosen.tech == tech].iloc[0]
        rows.append([tech, r["状态"], int(r["issue_pr_comments"]),
                     int(r["future_comments"]), "下降" if r["comments_down"] else "上升"])
    table(s, rows, 0.77, 1.52, [2.2, 3.4, 2.25, 2.25, 2.1], 0.68, 16)
    add_text(s, "Ruby 的问卷警报没有对应评论下降；Kotlin 只有评论警报却随后下降。",
             0.82, 5.87, 11.6, 0.7, 19, TEAL)

    s = slide(prs, "当前可用的核查清单", "依据：code/results/tables/multisignal_warning_v2_latest.csv；使用率只用于同级排序")
    rows = [["技术", "问卷", "评论", "输出", "使用率"]]
    for _, r in latest.iterrows():
        so_label = "不足" if not r["so_history_ok"] else ("警报" if r["so_alert"] else "无")
        gh_label = "警报" if r["gh_alert"] else "无"
        rows.append([r["tech"], so_label, gh_label, r["状态"], f"{100*r['usage']:.1f}%"])
    table(s, rows, 0.81, 1.43, [2.2, 1.55, 1.55, 4.65, 2.25], 0.58, 15)
    add_text(s, "Express、Flask 优先人工核查；Fastify、Phoenix 的问卷历史不足。",
             0.85, 6.27, 11.5, 0.42, 18, TEAL)

    s = slide(prs, "结论：先做核查排序，再扩大验证", "完整方法、逐项预测及限制见研究记录_多源因变量与动态预警.md")
    add_text(s, "加入评论量后，双源警报的命中比例提高，误报减少。",
             0.83, 1.54, 11.65, 0.76, 23, TEAL, True)
    add_text(s, "它同时漏掉更多下降；五项技术的样本不足以证明总体预测更强。",
             0.83, 2.79, 11.55, 1.0, 22, ORANGE)
    add_text(s, "实际用途：把双源收缩排在人工核查前面，单源冲突保留为待查。",
             0.83, 4.31, 11.55, 1.0, 21)
    add_text(s, "下一步：扩充独立仓库，区分机器人与人工评论，并检验新的未来年份。",
             0.83, 5.69, 11.5, 0.75, 20)

    OUT.parent.mkdir(exist_ok=True)
    prs.save(OUT)
    print(f"已生成 {OUT}（{len(prs.slides)} 页）")


if __name__ == "__main__":
    main()
